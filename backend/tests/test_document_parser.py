import os
import pytest
from docx import Document
from app.services.rag.document_parser import unified_document_parser
from app.services.rag.chunker import semantic_chunker


def test_docx_parsing_and_chunking(tmp_path):
    # 1. Create a mock Word document with heading, math formula and table
    docx_path = tmp_path / "test_physics_lesson.docx"
    doc = Document()
    doc.add_heading("高二物理《带电粒子在电场中的偏转》教学设计", level=1)
    doc.add_paragraph("本节课重点推导粒子离开偏转极板时的侧移距离：$y = \\frac{1}{2}at^2 = \\frac{qUL^2}{2mdv_0^2}$。")

    # Add a table
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "物理量"
    table.cell(0, 1).text = "符号"
    table.cell(0, 2).text = "单位"

    table.cell(1, 0).text = "偏转电场电压"
    table.cell(1, 1).text = "U"
    table.cell(1, 2).text = "V"

    doc.save(str(docx_path))

    # 2. Parse using unified_document_parser
    parsed = unified_document_parser.parse_file(str(docx_path))

    assert parsed["doc_type"] == "WORD"
    assert "高二物理《带电粒子在电场中的偏转》" in parsed["markdown_content"]
    assert "qUL^2" in parsed["markdown_content"]
    # Check that table is converted to markdown table
    assert "| 物理量 | 符号 | 单位 |" in parsed["markdown_content"]
    assert ("|---" in parsed["markdown_content"] or "| ---" in parsed["markdown_content"])

    assert parsed["table_count"] >= 1
    assert parsed["formula_count"] >= 1

    # 3. Test chunking on parsed markdown
    chunks = semantic_chunker.chunk_document(
        markdown_text=parsed["markdown_content"],
        doc_id="test_doc_word",
        kb_id="kb_physics",
        default_metadata={"subject": "高中物理"}
    )
    assert len(chunks) >= 1
    assert chunks[0]["metadata"]["has_formula"] is True
    assert chunks[0]["metadata"]["has_table"] is True
