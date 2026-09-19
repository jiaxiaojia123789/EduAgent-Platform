import io
import os
import re
import urllib.parse

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

router = APIRouter(prefix="/artifacts", tags=["Generated Artifacts & Export"])


# ============================================================
# 共享：markdown 行内解析 + LaTeX→Unicode 映射
# Word 与 PDF 导出共用，保证两种格式输出一致
# ============================================================

# Windows 系统中文字体候选路径（按优先级）
_FONT_CANDIDATES = {
    "regular": [
        r"C:\Windows\Fonts\msyh.ttc",    # 微软雅黑
        r"C:\Windows\Fonts\simhei.ttf",  # 黑体
        r"C:\Windows\Fonts\simsun.ttc",  # 宋体
    ],
    "bold": [
        r"C:\Windows\Fonts\msyhbd.ttc",  # 微软雅黑粗体
        r"C:\Windows\Fonts\simhei.ttf",  # 黑体（兼作粗体）
    ],
}

_WORD_FONT = "微软雅黑"  # docx 中的中文字体名

# 上标 / 下标字符表（数字 + 常见字母）
_SUPER_MAP = str.maketrans("0123456789+-=()ni", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿⁱ")
_SUB_MAP = str.maketrans("0123456789+-=()nijmh", "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎ᵢⱼₙₘₕ")

# 常见 LaTeX 命令 → Unicode 符号
_LATEX_SYMBOLS = {
    "alpha": "α", "beta": "β", "gamma": "γ", "delta": "δ", "epsilon": "ε",
    "theta": "θ", "lambda": "λ", "mu": "μ", "nu": "ν", "pi": "π", "rho": "ρ",
    "sigma": "σ", "tau": "τ", "phi": "φ", "omega": "ω",
    "Delta": "Δ", "Sigma": "Σ", "Omega": "Ω", "Theta": "Θ", "Phi": "Φ",
    "times": "×", "div": "÷", "pm": "±", "mp": "∓", "cdot": "·",
    "leq": "≤", "le": "≤", "geq": "≥", "ge": "≥", "neq": "≠", "ne": "≠",
    "approx": "≈", "equiv": "≡", "sim": "∼", "propto": "∝",
    "infty": "∞", "sum": "∑", "prod": "∏", "int": "∫",
    "rightarrow": "→", "to": "→", "leftarrow": "←", "Rightarrow": "⇒",
    "in": "∈", "subset": "⊂", "cup": "∪", "cap": "∩",
    "angle": "∠", "parallel": "∥", "perp": "⊥", "degree": "°",
    "because": "∵", "therefore": "∴", "prime": "′",
}


def _to_super(s: str) -> str:
    # translate 对映射表外的字符原样保留
    return s.translate(_SUPER_MAP)


def _to_sub(s: str) -> str:
    return s.translate(_SUB_MAP)


def _latex_to_unicode(s: str) -> str:
    """将轻量 LaTeX 数学片段转换为可读 Unicode 文本，未识别命令原样清理。"""
    # 结构性命令优先：根号、分式
    s = re.sub(r"\\sqrt\{([^{}]*)\}", r"√(\1)", s)
    s = re.sub(r"\\frac\{([^{}]*)\}\{([^{}]*)\}", r"(\1)/(\2)", s)
    # 带花括号的上下标
    s = re.sub(r"\^\{([^{}]*)\}", lambda m: _to_super(m.group(1)), s)
    s = re.sub(r"_\{([^{}]*)\}", lambda m: _to_sub(m.group(1)), s)
    # 单字符上下标（x^2、v_0）
    s = re.sub(r"\^([0-9a-zA-Z])", lambda m: _to_super(m.group(1)), s)
    s = re.sub(r"_([0-9a-zA-Z])", lambda m: _to_sub(m.group(1)), s)
    # 希腊字母与运算符
    for cmd, rep in _LATEX_SYMBOLS.items():
        s = re.sub(r"\\" + cmd + r"(?![a-zA-Z])", rep, s)
    # 数学函数命令：去反斜杠保留原文（\cos → cos）
    for fn in ("arcsin", "arccos", "arctan", "sin", "cos", "tan", "cot", "sec", "csc",
               "log", "ln", "lg", "exp", "max", "min", "vec", "hat"):
        s = re.sub(r"\\" + fn + r"(?![a-zA-Z])", fn, s)
    # 清理残留：\left \right、空白命令、未知命令、定界符
    s = s.replace("\\left", "").replace("\\right", "")
    s = re.sub(r"\\[,;:! ]", " ", s)
    s = re.sub(r"\\[a-zA-Z]+", "", s)
    s = s.replace("$", "")
    return s


# 行内标记解析：**粗体** / *斜体* / `代码` / $行内公式$（** 必须在 * 之前）
_INLINE_RE = re.compile(
    r"(\*\*[^*\n]+?\*\*)"
    r"|(\*[^*\n]+?\*)"
    r"|(`[^`\n]+?`)"
    r"|(\$[^$\n]+?\$)"
)


def _parse_inline(text: str) -> list[tuple[str, str]]:
    """把一行文本拆成 (内容, 类型) 片段序列，类型：plain/bold/italic/code/math。"""
    parts: list[tuple[str, str]] = []
    pos = 0
    for m in _INLINE_RE.finditer(text):
        if m.start() > pos:
            parts.append((text[pos:m.start()], "plain"))
        tok = m.group(0)
        if tok.startswith("**"):
            parts.append((tok[2:-2], "bold"))
        elif tok.startswith("`"):
            parts.append((tok[1:-1], "code"))
        elif tok.startswith("$"):
            parts.append((_latex_to_unicode(tok[1:-1]), "math"))
        else:
            parts.append((tok[1:-1], "italic"))
        pos = m.end()
    if pos < len(text):
        parts.append((text[pos:], "plain"))
    return parts


def _clean_inline(text: str) -> str:
    """表格单元格等纯文本场景：去掉行内标记符号并映射公式。"""
    return "".join(
        _latex_to_unicode(seg) if kind == "math" else seg
        for seg, kind in _parse_inline(text)
    ).replace("**", "").replace("`", "")


def _resolve_font(candidates: list[str]) -> str | None:
    """返回第一个真实存在的字体文件路径。"""
    for path in candidates:
        if os.path.exists(path):
            return path
    return None


def _split_table_row(line: str) -> list[str]:
    """把 | a | b | 行拆成单元格列表。"""
    cells = line.strip().strip("|").split("|")
    return [c.strip() for c in cells]


class ExportRequest(BaseModel):
    title: str
    markdown_content: str
    artifact_type: str = "LESSON_PLAN"


# ============================================================
# Word 导出
# ============================================================

def _set_run_font(run, name: str = _WORD_FONT):
    """同时设置西文与中文字体，避免中文回退宋体。"""
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)


def _add_rich_runs(p, text: str, *, size: Pt | None = None, color: RGBColor | None = None, base_bold: bool = False):
    """按行内解析结果向段落添加富文本 runs。"""
    for seg, kind in _parse_inline(text):
        if not seg:
            continue
        r = p.add_run(seg)
        if size:
            r.font.size = size
        if color:
            r.font.color.rgb = color
        if base_bold or kind == "bold":
            r.font.bold = True
        if kind == "italic":
            r.font.italic = True
        if kind == "code":
            r.font.name = "Consolas"
            r.font.color.rgb = RGBColor(190, 24, 93)
        elif kind == "math":
            r.font.color.rgb = RGBColor(15, 23, 42)
        _set_run_font(r, "Consolas" if kind == "code" else _WORD_FONT)
    return p


def _add_word_hr(doc):
    """插入真正的水平分割线（段落底边框）。"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "CBD5E1")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _add_word_quote(doc, quote_lines: list[str]):
    """渲染引用块：左侧缩进 + 灰色文字 + 逐行富文本。"""
    for q in quote_lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        _add_rich_runs(p, q, size=Pt(11), color=RGBColor(71, 85, 105))
    return doc


def _add_word_table(doc, rows: list[list[str]]):
    """渲染 markdown 表格（首行为表头）。"""
    cols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=cols)
    t.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci in range(cols):
            cell = t.cell(ri, ci)
            text = _clean_inline(row[ci]) if ci < len(row) else ""
            cell.text = ""
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(10)
            if ri == 0:
                r.font.bold = True
            _set_run_font(r)


@router.post("/export/word")
async def export_to_word(payload: ExportRequest):
    """
    Exports structured lesson plan or exam paper as a styled Word document (.docx).
    支持：标题层级、加粗/斜体/行内代码、行内 LaTeX(映射为 Unicode)、
    引用块、无序列表、分割线、markdown 表格、$$ 公式块。
    """
    doc = Document()

    # Set Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(payload.title)
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(30, 58, 138)
    _set_run_font(title_run)

    # Add subtitle
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("智学中台 (EduAgent-Platform) · 智能生成教学成果")
    sub_run.font.size = Pt(10)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(100, 116, 139)
    _set_run_font(sub_run)

    doc.add_paragraph()  # Blank line

    lines = payload.markdown_content.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines):
        line_s = lines[i].strip()
        if not line_s:
            i += 1
            continue

        # 1. markdown 表格（当前行 | 开头且下一行为 |---| 分隔行）
        if line_s.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:\-|]+\|$", lines[i + 1].strip()):
            rows = [_split_table_row(line_s)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_split_table_row(lines[i].strip()))
                i += 1
            _add_word_table(doc, rows)
            continue

        # 2. $$ 公式块（LaTeX 映射后居中）
        if line_s.startswith("$$") and line_s.endswith("$$") and len(line_s) > 4:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(_latex_to_unicode(line_s.strip("$").strip()))
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(15, 23, 42)
            _set_run_font(r)
            i += 1
            continue

        # 3. 标题层级
        if line_s.startswith("### "):
            h = doc.add_heading(level=2)
            run = h.add_run(line_s[4:])
            run.font.size = Pt(14)
            run.font.bold = True
            _set_run_font(run)
            i += 1
            continue
        if line_s.startswith("## "):
            h = doc.add_heading(level=1)
            run = h.add_run(line_s[3:])
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138)
            _set_run_font(run)
            i += 1
            continue
        if line_s.startswith("# "):
            h = doc.add_heading(level=1)
            run = h.add_run(line_s[2:])
            run.font.size = Pt(18)
            run.font.bold = True
            _set_run_font(run)
            i += 1
            continue

        # 4. 引用块（连续 > 行聚合）
        if line_s.startswith(">"):
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                q = lines[i].strip()[1:].strip()
                if q:
                    quote_lines.append(q)
                i += 1
            _add_word_quote(doc, quote_lines)
            continue

        # 5. 分割线
        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", line_s):
            _add_word_hr(doc)
            i += 1
            continue

        # 6. 无序列表
        if line_s.startswith("- ") or line_s.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            _add_rich_runs(p, line_s[2:], size=Pt(11), color=RGBColor(30, 41, 59))
            i += 1
            continue

        # 7. 普通段落（行内富文本解析）
        p = doc.add_paragraph()
        _add_rich_runs(p, line_s, size=Pt(11), color=RGBColor(30, 41, 59))
        i += 1

    # Save to memory buffer
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)

    filename = f"{payload.title.replace(' ', '_')}.docx"
    encoded_filename = urllib.parse.quote(filename)
    return StreamingResponse(
        file_stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"},
    )


# ============================================================
# PDF 导出
# ============================================================

_COLOR_TITLE = (30, 58, 138)
_COLOR_SUB = (100, 116, 139)
_COLOR_H1 = (15, 23, 42)
_COLOR_H2 = (30, 58, 138)
_COLOR_BODY = (30, 41, 59)
_COLOR_QUOTE = (71, 85, 105)
_COLOR_CODE = (190, 24, 93)
_COLOR_LINE = (203, 213, 225)


@router.post("/export/pdf")
async def export_to_pdf(payload: ExportRequest):
    """
    Exports structured lesson plan or exam paper as a styled PDF document.
    Uses Windows system Chinese fonts (Microsoft YaHei preferred) for CJK support.
    支持：标题层级、加粗/斜体/行内代码、行内 LaTeX(映射为 Unicode)、
    引用块、无序列表、分割线、markdown 表格、$$ 公式块。
    """
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    font_regular = _resolve_font(_FONT_CANDIDATES["regular"])
    font_bold = _resolve_font(_FONT_CANDIDATES["bold"]) or font_regular
    if not font_regular:
        raise HTTPException(status_code=500, detail="未找到可用的中文字体文件，无法导出 PDF")

    pdf = FPDF(format="A4")
    pdf.set_margins(20, 18, 20)
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()

    # 注册中文字体（regular / bold 两种样式）
    pdf.add_font("cjk", "", font_regular)
    if font_bold != font_regular:
        pdf.add_font("cjk", "B", font_bold)
    else:
        # 黑体本身笔画较粗，直接映射 B 样式
        pdf.add_font("cjk", "B", font_regular)

    # 上下标等符号（₀¹ⁿ）在中文字体中缺字形，注册 Segoe UI 作为兜底
    segoe = r"C:\Windows\Fonts\segoeui.ttf"
    if os.path.exists(segoe):
        pdf.add_font("fallback", "", segoe)
        pdf.set_fallback_fonts(["fallback"])

    # multi_cell 渲染后光标回到左边距并换到下一行
    _back = {"new_x": XPos.LMARGIN, "new_y": YPos.NEXT}

    def write_rich(text: str, size: float = 11, color: tuple = _COLOR_BODY, indent: float = 0):
        """逐段切换字体写出富文本（write 支持自动换行），行尾换行。"""
        if indent:
            pdf.set_x(pdf.l_margin + indent)
        line_h = size * 0.58
        for seg, kind in _parse_inline(text):
            if not seg:
                continue
            if kind == "bold":
                pdf.set_font("cjk", "B", size)
            else:
                pdf.set_font("cjk", "", size)
            seg_color = color
            if kind == "code":
                seg_color = _COLOR_CODE
            elif kind == "math":
                seg_color = _COLOR_H1
            pdf.set_text_color(*seg_color)
            pdf.write(line_h, seg)
        pdf.ln(line_h + 1)

    # 主标题（居中、深蓝、粗体）——与 Word 导出风格保持一致
    pdf.set_font("cjk", "B", 20)
    pdf.set_text_color(*_COLOR_TITLE)
    pdf.multi_cell(0, 12, payload.title, align="C", **_back)

    # 副标题
    pdf.set_font("cjk", "", 10)
    pdf.set_text_color(*_COLOR_SUB)
    pdf.multi_cell(0, 7, "智学中台 (EduAgent-Platform) · 智能生成教学成果", align="C", **_back)

    pdf.ln(6)

    lines = payload.markdown_content.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines):
        line_s = lines[i].strip()
        if not line_s:
            pdf.ln(2)
            i += 1
            continue

        # 1. markdown 表格
        if line_s.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:\-|]+\|$", lines[i + 1].strip()):
            rows = [_split_table_row(line_s)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_split_table_row(lines[i].strip()))
                i += 1
            pdf.set_font("cjk", "", 10)
            pdf.set_text_color(*_COLOR_BODY)
            pdf.set_draw_color(148, 163, 184)
            with pdf.table(line_height=6, padding=2, text_align="LEFT") as table:
                for ri, row in enumerate(rows):
                    trow = table.row()
                    for cell_text in row:
                        trow.cell(_clean_inline(cell_text))
            pdf.ln(2)
            continue

        # 2. $$ 公式块（LaTeX 映射后居中）
        if line_s.startswith("$$") and line_s.endswith("$$") and len(line_s) > 4:
            pdf.set_font("cjk", "", 11)
            pdf.set_text_color(*_COLOR_H1)
            pdf.multi_cell(0, 7, _latex_to_unicode(line_s.strip("$").strip()), align="C", **_back)
            pdf.ln(1)
            i += 1
            continue

        # 3. 标题层级
        if line_s.startswith("### "):
            pdf.set_font("cjk", "B", 14)
            pdf.set_text_color(*_COLOR_H1)
            pdf.multi_cell(0, 8, line_s[4:], **_back)
            pdf.ln(1)
            i += 1
            continue
        if line_s.startswith("## "):
            pdf.ln(2)
            pdf.set_font("cjk", "B", 16)
            pdf.set_text_color(*_COLOR_H2)
            pdf.multi_cell(0, 9, line_s[3:], **_back)
            pdf.ln(1)
            i += 1
            continue
        if line_s.startswith("# "):
            pdf.ln(2)
            pdf.set_font("cjk", "B", 18)
            pdf.set_text_color(*_COLOR_H1)
            pdf.multi_cell(0, 10, line_s[2:], **_back)
            pdf.ln(1)
            i += 1
            continue

        # 4. 引用块（▎竖线前缀 + 灰色富文本）
        if line_s.startswith(">"):
            while i < len(lines) and lines[i].strip().startswith(">"):
                q = lines[i].strip()[1:].strip()
                if q:
                    write_rich("▎" + q, size=10.5, color=_COLOR_QUOTE, indent=2)
                i += 1
            pdf.ln(1)
            continue

        # 5. 分割线
        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", line_s):
            pdf.ln(2)
            pdf.set_draw_color(*_COLOR_LINE)
            pdf.line(pdf.l_margin, pdf.get_y(), pdf.w - pdf.r_margin, pdf.get_y())
            pdf.ln(4)
            i += 1
            continue

        # 6. 无序列表
        if line_s.startswith("- ") or line_s.startswith("* "):
            write_rich("• " + line_s[2:], size=11, color=_COLOR_BODY, indent=4)
            i += 1
            continue

        # 7. 普通段落（行内富文本解析）
        write_rich(line_s, size=11, color=_COLOR_BODY)
        pdf.ln(1)
        i += 1

    file_stream = io.BytesIO()
    pdf.output(file_stream)
    file_stream.seek(0)

    filename = f"{payload.title.replace(' ', '_')}.pdf"
    encoded_filename = urllib.parse.quote(filename)
    return StreamingResponse(
        file_stream,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"},
    )
